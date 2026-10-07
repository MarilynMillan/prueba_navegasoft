# -*- coding: utf-8 -*-
import base64
import io
import logging
from odoo import fields, models, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class DSInventoryAdjustment(models.Model):
    """Ajuste de Inventario Retroactivo — modelo unificado con estados.

    Borrador: el usuario puede editar cantidades contadas, importar, recalcular.
    Aplicado: se generaron los stock.move y queda como registro histórico readonly.
    """
    _name = 'ds.inventory.adjustment'
    _description = 'Ajuste de Inventario Retroactivo'
    _order = 'create_date desc, id desc'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(
        string='Referencia',
        required=True,
        readonly=True,
        default='Nuevo',
        copy=False,
    )
    state = fields.Selection([
        ('draft', 'Borrador'),
        ('done', 'Aplicado'),
    ], default='draft', string='Estado', tracking=True)

    adjustment_datetime = fields.Datetime(
        string='Fecha y Hora del Conteo',
        required=True,
        tracking=True,
    )
    location_id = fields.Many2one(
        'stock.location',
        string='Ubicación',
        required=True,
        domain="[('usage', '=', 'internal')]",
        tracking=True,
    )
    company_id = fields.Many2one(
        'res.company',
        string='Compañía',
        required=True,
        readonly=True,
        default=lambda self: self.env.company,
    )
    currency_id = fields.Many2one(
        'res.currency',
        string='Moneda',
        related='company_id.currency_id',
        store=True,
    )

    # ---------- Filtros de productos ----------
    product_filter = fields.Selection([
        ('all', 'Todos los productos'),
        ('category', 'Por categoría'),
        ('products', 'Productos específicos'),
    ], string='Filtro de Productos', default='all', required=True)
    category_ids = fields.Many2many(
        'product.category',
        'ds_inv_adj_categ_rel',
        'adjustment_id',
        'category_id',
        string='Categorías',
    )
    filter_product_ids = fields.Many2many(
        'product.product',
        'ds_inv_adj_filter_product_rel',
        'adjustment_id',
        'product_id',
        string='Productos',
        domain="[('is_storable', '=', True)]",
    )

    # ---------- Datos de confirmación ----------
    reason = fields.Text(
        string='Motivo del Inventario',
        readonly=True,
    )
    responsible_employee_ids = fields.Many2many(
        'hr.employee',
        'ds_inv_adj_employee_rel',
        'adjustment_id',
        'employee_id',
        string='Responsables del Conteo',
        readonly=True,
    )
    applied_datetime = fields.Datetime(
        string='Fecha de Aplicación',
        readonly=True,
    )

    # ---------- Líneas y movimientos ----------
    line_ids = fields.One2many(
        'ds.inventory.adjustment.line',
        'adjustment_id',
        string='Líneas de Ajuste',
    )
    move_ids = fields.One2many(
        'stock.move',
        'ds_adjustment_record_id',
        string='Movimientos de Stock',
        readonly=True,
    )

    # ---------- Campos snapshot ----------
    total_inventory_value = fields.Monetary(
        string='Valoración Total Inventario',
        currency_field='currency_id',
        readonly=True,
    )
    total_adjustment_value = fields.Monetary(
        string='Valor Total de Ajustes',
        currency_field='currency_id',
        readonly=True,
    )
    line_count = fields.Integer(
        string='Productos',
        compute='_compute_line_count',
    )
    move_count = fields.Integer(
        string='Movimientos',
        compute='_compute_move_count',
    )
    is_calculated = fields.Boolean(
        string='Calculado',
        default=False,
        help='Indica si ya se calcularon las cantidades teóricas.',
    )

    def _compute_line_count(self):
        for record in self:
            record.line_count = len(record.line_ids)

    def _compute_move_count(self):
        for record in self:
            record.move_count = len(record.move_ids)

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', 'Nuevo') == 'Nuevo':
                vals['name'] = self.env['ir.sequence'].next_by_code(
                    'ds.inventory.adjustment'
                ) or 'Nuevo'
        records = super().create(vals_list)
        # Advertencia de duplicado
        for record in records:
            record._check_duplicate_warning()
        return records

    def _check_duplicate_warning(self):
        """Advierte si ya existe un borrador con misma ubicación, fecha y hora."""
        self.ensure_one()
        if not self.adjustment_datetime or not self.location_id:
            return
        duplicate = self.search([
            ('id', '!=', self.id),
            ('state', '=', 'draft'),
            ('location_id', '=', self.location_id.id),
            ('adjustment_datetime', '=', self.adjustment_datetime),
        ], limit=1)
        if duplicate:
            self.message_post(
                body=_(
                    '<strong>⚠️ Advertencia:</strong> Ya existe un borrador '
                    '<b>%(name)s</b> para la misma ubicación, fecha y hora. '
                    'Verifique que no esté duplicando el trabajo.',
                    name=duplicate.name,
                ),
                subtype_xmlid='mail.mt_note',
            )

    # =====================================================================
    # CALCULAR / RECALCULAR
    # =====================================================================
    def action_calculate(self):
        """Calcula la cantidad teórica a la fecha/hora seleccionada."""
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se puede calcular en estado Borrador.'))
        if self.adjustment_datetime > fields.Datetime.now():
            raise UserError(_('La fecha y hora del conteo no puede ser en el futuro.'))

        location_id = self.location_id.id
        cutoff_dt = self.adjustment_datetime

        # Paso 1: Quants actuales con filtros
        quant_domain = [
            ('location_id', '=', location_id),
            ('product_id.is_storable', '=', True),
            ('quantity', '!=', 0),
        ]
        if self.product_filter == 'category' and self.category_ids:
            quant_domain.append(('product_id.categ_id', 'child_of', self.category_ids.ids))
        elif self.product_filter == 'products' and self.filter_product_ids:
            quant_domain.append(('product_id', 'in', self.filter_product_ids.ids))

        quants = self.env['stock.quant'].sudo().search(quant_domain)
        if not quants and not self.line_ids:
            raise UserError(_(
                'No se encontraron productos con stock en la ubicación seleccionada '
                'con los filtros aplicados.'
            ))

        product_ids = quants.mapped('product_id').ids
        current_qty_map = {}
        for q in quants:
            current_qty_map.setdefault(q.product_id.id, 0.0)
            current_qty_map[q.product_id.id] += q.quantity

        # También incluir productos de líneas existentes en la query SQL
        existing_product_ids = self.line_ids.mapped('product_id').ids
        all_product_ids = list(set(product_ids + existing_product_ids))

        # Paso 2: SQL para movimientos posteriores
        net_after_map = {}
        if all_product_ids:
            query = """
                SELECT
                    sml.product_id,
                    COALESCE(SUM(
                        CASE WHEN sml.location_dest_id = %(loc)s THEN sml.quantity ELSE 0 END
                    ), 0)
                    -
                    COALESCE(SUM(
                        CASE WHEN sml.location_id = %(loc)s THEN sml.quantity ELSE 0 END
                    ), 0) AS net_after
                FROM stock_move_line sml
                INNER JOIN stock_move sm ON sm.id = sml.move_id
                WHERE sm.state = 'done'
                  AND sml.date > %(cutoff)s
                  AND sml.product_id = ANY(%(products)s)
                  AND (sml.location_id = %(loc)s OR sml.location_dest_id = %(loc)s)
                GROUP BY sml.product_id
            """
            self.env.cr.execute(query, {
                'loc': location_id,
                'cutoff': cutoff_dt,
                'products': all_product_ids,
            })
            net_after_map = {row[0]: row[1] for row in self.env.cr.fetchall()}

        # Paso 3: Crear o actualizar líneas, NUNCA eliminar
        existing_lines = {l.product_id.id: l for l in self.line_ids}
        line_vals = []
        Product = self.env['product.product'].sudo()
        all_products = Product.browse(all_product_ids)
        product_map = {p.id: p for p in all_products}

        for pid, current_qty in current_qty_map.items():
            product = product_map.get(pid)
            if not product:
                continue
            net_after = net_after_map.get(pid, 0.0)
            theoretical_qty = current_qty - net_after

            if pid in existing_lines:
                existing_lines[pid].theoretical_qty = theoretical_qty
            else:
                line_vals.append({
                    'adjustment_id': self.id,
                    'product_id': pid,
                    'product_uom_id': product.uom_id.id,
                    'theoretical_qty': theoretical_qty,
                    'counted_qty': 0.0,
                })

        # Líneas existentes sin quant actual: recalcular teórica
        for pid, line in existing_lines.items():
            if pid not in current_qty_map:
                net_after = net_after_map.get(pid, 0.0)
                line.theoretical_qty = 0.0 - net_after

        if line_vals:
            self.env['ds.inventory.adjustment.line'].create(line_vals)

        self.is_calculated = True

    # =====================================================================
    # IMPORTAR CONTEO
    # =====================================================================
    def action_open_import_wizard(self):
        """Abre wizard de importación."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importar Conteo',
            'res_model': 'ds.inventory.adjustment.import',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_adjustment_id': self.id},
        }

    # =====================================================================
    # APLICAR AJUSTE
    # =====================================================================
    def action_open_confirm_dialog(self):
        """Abre diálogo de confirmación con motivo y responsables."""
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Este ajuste ya fue aplicado.'))

        lines_with_diff = self.line_ids.filtered(lambda l: l.difference_qty != 0)
        if not lines_with_diff:
            raise UserError(_('No hay diferencias para ajustar. Verifique las cantidades contadas.'))

        return {
            'type': 'ir.actions.act_window',
            'name': 'Confirmar Ajuste Retroactivo',
            'res_model': 'ds.inventory.adjustment.confirm',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_adjustment_id': self.id},
        }

    def apply_adjustment(self, reason, employee_ids):
        """Aplica el ajuste retroactivo. Llamado desde el wizard de confirmación."""
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Este ajuste ya fue aplicado.'))

        lines_to_adjust = self.line_ids.filtered(lambda l: l.difference_qty != 0)
        if not lines_to_adjust:
            raise UserError(_('No hay diferencias para ajustar.'))

        backdate = self.adjustment_datetime
        location = self.location_id
        company = self.company_id

        # Crear stock.moves en batch — idéntico al nativo _apply_inventory
        move_vals_list = []
        for line in lines_to_adjust:
            diff = line.difference_qty
            product = line.product_id
            inventory_location = product.with_company(company).property_stock_inventory

            if diff > 0:
                src_location = inventory_location
                dst_location = location
            else:
                src_location = location
                dst_location = inventory_location

            move_vals_list.append({
                'name': _('Ajuste Retroactivo: %(product)s', product=product.display_name),
                'product_id': product.id,
                'product_uom_qty': abs(diff),
                'product_uom': line.product_uom_id.id,
                'location_id': src_location.id,
                'location_dest_id': dst_location.id,
                'company_id': company.id,
                'state': 'confirmed',
                'is_inventory': True,
                'picked': True,
                'ds_adjustment_record_id': self.id,
                'move_line_ids': [(0, 0, {
                    'product_id': product.id,
                    'product_uom_id': line.product_uom_id.id,
                    'quantity': abs(diff),
                    'location_id': src_location.id,
                    'location_dest_id': dst_location.id,
                    'company_id': company.id,
                })],
            })

        moves = self.env['stock.move'].with_context(
            inventory_mode=False
        ).sudo().create(move_vals_list)

        # Validar — flujo idéntico al nativo
        moves._action_done()

        # Aplicar backdate
        self._apply_backdate(moves, backdate)

        # Calcular snapshots de valoración
        total_inventory_value = 0.0
        total_adjustment_value = 0.0
        for line in lines_to_adjust:
            unit_cost = line.product_id.standard_price
            line.unit_cost = unit_cost
            line.difference_value = line.difference_qty * unit_cost
            total_inventory_value += line.counted_qty * unit_cost
            total_adjustment_value += line.difference_value

        # Actualizar registro
        self.write({
            'state': 'done',
            'reason': reason,
            'responsible_employee_ids': [(6, 0, employee_ids)],
            'applied_datetime': fields.Datetime.now(),
            'total_inventory_value': total_inventory_value,
            'total_adjustment_value': total_adjustment_value,
        })

        # Log en chatter
        employee_names = ', '.join(
            self.env['hr.employee'].browse(employee_ids).mapped('name')
        )
        self.message_post(
            body=_(
                '<strong>Ajuste de Inventario Retroactivo aplicado</strong><br/>'
                '<b>Fecha del conteo:</b> %(date)s<br/>'
                '<b>Ubicación:</b> %(location)s<br/>'
                '<b>Motivo:</b> %(reason)s<br/>'
                '<b>Responsables:</b> %(employees)s<br/>'
                '<b>Productos ajustados:</b> %(count)s<br/>'
                '<b>Valor neto del ajuste:</b> %(value)s',
                date=str(backdate),
                location=location.complete_name,
                reason=reason,
                employees=employee_names,
                count=len(lines_to_adjust),
                value=total_adjustment_value,
            ),
            subtype_xmlid='mail.mt_note',
        )

    def _apply_backdate(self, moves, backdate):
        """Aplica fecha retroactiva a moves, move_lines, account_moves y SVL."""
        moves.write({'date': backdate})
        moves.move_line_ids.write({'date': backdate})

        account_moves = self.env['account.move'].sudo().search([
            ('stock_move_id', 'in', moves.ids),
        ])
        if account_moves:
            account_moves.button_draft()
            account_moves.write({'name': False, 'date': backdate.date()})
            account_moves.action_post()

        if moves.ids:
            self.env.cr.execute("""
                UPDATE stock_valuation_layer
                SET create_date = %(backdate)s
                WHERE stock_move_id IN %(move_ids)s
            """, {'backdate': backdate, 'move_ids': tuple(moves.ids)})

    # =====================================================================
    # SMART BUTTONS Y ACCIONES
    # =====================================================================
    def action_view_moves(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': 'Movimientos del Ajuste',
            'res_model': 'stock.move',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.move_ids.ids)],
            'context': {'create': False},
        }

    def action_export_xlsx(self):
        """Exporta el detalle del ajuste como archivo Excel."""
        self.ensure_one()
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            raise UserError('La librería openpyxl no está disponible.')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ajuste Retroactivo'

        header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        title_font = Font(name='Arial', bold=True, size=14)
        info_font = Font(name='Arial', size=10)
        info_bold = Font(name='Arial', bold=True, size=10)
        number_format = '#,##0.0000'
        money_format = '#,##0.00'
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        ws.merge_cells('A1:H1')
        ws['A1'] = f'Ajuste de Inventario Retroactivo - {self.name}'
        ws['A1'].font = title_font

        info_rows = [
            ('Fecha del Conteo:', str(self.adjustment_datetime)),
            ('Ubicación:', self.location_id.complete_name),
            ('Motivo:', self.reason or ''),
            ('Responsables:', ', '.join(self.responsible_employee_ids.mapped('name'))),
            ('Fecha de Aplicación:', str(self.applied_datetime or '')),
            ('Valor Total Inventario:', f'{self.total_inventory_value:,.2f}'),
            ('Valor Total Ajustes:', f'{self.total_adjustment_value:,.2f}'),
        ]
        for i, (label, value) in enumerate(info_rows, start=3):
            ws[f'A{i}'] = label
            ws[f'A{i}'].font = info_bold
            ws[f'B{i}'] = value
            ws[f'B{i}'].font = info_font

        header_row = 12
        headers = ['Ref. Interna', 'Producto', 'UdM', 'Teórica a Fecha',
                   'Contada', 'Diferencia', 'Costo Unit.', 'Valor Diferencia']
        col_widths = [15, 40, 10, 18, 15, 15, 15, 18]

        for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
            ws.column_dimensions[cell.column_letter].width = width

        for row_idx, line in enumerate(self.line_ids, start=header_row + 1):
            data = [
                line.product_id.default_code or '',
                line.product_id.display_name,
                line.product_uom_id.name,
                line.theoretical_qty, line.counted_qty, line.difference_qty,
                line.unit_cost, line.difference_value,
            ]
            for col_idx, value in enumerate(data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.font = Font(name='Arial', size=10)
                if col_idx >= 4:
                    cell.number_format = number_format if col_idx <= 6 else money_format
                    cell.alignment = Alignment(horizontal='right')

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_data = base64.b64encode(output.read())
        output.close()

        filename = f'Ajuste_Retroactivo_{self.name.replace("/", "_")}.xlsx'
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': file_data,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class DSInventoryAdjustmentLine(models.Model):
    """Línea de ajuste — persistente para soportar borradores."""
    _name = 'ds.inventory.adjustment.line'
    _description = 'Línea de Ajuste de Inventario Retroactivo'
    _order = 'product_id'

    adjustment_id = fields.Many2one(
        'ds.inventory.adjustment',
        string='Ajuste',
        required=True,
        ondelete='cascade',
        index=True,
    )
    state = fields.Selection(
        related='adjustment_id.state',
        store=True,
    )
    currency_id = fields.Many2one(
        'res.currency',
        related='adjustment_id.currency_id',
        store=True,
    )
    product_id = fields.Many2one(
        'product.product',
        string='Producto',
        required=True,
        readonly=True,
    )
    product_uom_id = fields.Many2one(
        'uom.uom',
        string='UdM',
        readonly=True,
    )
    theoretical_qty = fields.Float(
        string='Cantidad Teórica a Fecha',
        digits='Product Unit of Measure',
        readonly=True,
    )
    counted_qty = fields.Float(
        string='Cantidad Contada',
        digits='Product Unit of Measure',
    )
    difference_qty = fields.Float(
        string='Diferencia',
        digits='Product Unit of Measure',
        compute='_compute_difference',
        store=True,
    )
    unit_cost = fields.Monetary(
        string='Costo Unitario',
        currency_field='currency_id',
        readonly=True,
    )
    difference_value = fields.Monetary(
        string='Valor de la Diferencia',
        currency_field='currency_id',
        readonly=True,
    )

    @api.depends('counted_qty', 'theoretical_qty')
    def _compute_difference(self):
        for line in self:
            line.difference_qty = line.counted_qty - line.theoretical_qty
